"""Pitch deck: general overview first, then workflow, RL/DQN, randomness."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = 13.333333, 7.5
NAVY = "122B45"
CREAM = "F6F4EF"
COPPER = "C45C26"
INK = "1C2430"
MUTED = "5C6B7A"
WHITE = "FFFFFF"
LINE = "E2DCD2"
WARM = "E8B896"
SOFT = "E6E1D8"
PALE = "F3EFE8"
FONT = "Calibri"
TOTAL = 20
OUT = "/home/waqas-ahmad/Desktop/reddit/docs/Reddit_CRM_tool.pptx"


def rgb(value):
    return RGBColor.from_string(value)


def style_run(run, size, color, bold=False, spacing=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    if spacing:
        rPr.set("spc", str(spacing))
    for tag in (qn("a:latin"), qn("a:ea"), qn("a:cs")):
        for child in list(rPr):
            if child.tag == tag:
                rPr.remove(child)
    for tag in (qn("a:latin"), qn("a:ea"), qn("a:cs")):
        node = etree.SubElement(rPr, tag)
        node.set("typeface", FONT)


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, attr, Emu(0))
    return tf


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, before=0, after=0, spacing=1.0, first=False, spc=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.line_spacing = spacing
    run = p.add_run()
    run.text = text
    style_run(run, size, color, bold, spc)
    return p


def rect(slide, x, y, w, h, fill, line=None, rounded=False, radius=0.08):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.0)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = radius
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def content_bg(slide):
    rect(slide, 0, 0, W, H, CREAM)
    rect(slide, 0, 0, 0.09, H, NAVY)


def footer(slide, page, dark=False):
    color = "8FA0B0" if dark else MUTED
    tf = box(slide, 0.48, 7.1, 7.2, 0.26)
    para(tf, "Klarivo   ·   Internal", 11, color, first=True)
    tf = box(slide, 9.4, 7.1, 3.45, 0.26)
    para(tf, f"{page}   /   {TOTAL}", 11, color, align=PP_ALIGN.RIGHT, first=True)


def eyebrow(slide, text, y=0.26):
    tf = box(slide, 0.48, y, 12.2, 0.26)
    para(tf, text, 12, COPPER, bold=True, first=True, spc=120)


def title(slide, text, y=0.5, h=0.46, size=28):
    tf = box(slide, 0.48, y, 12.3, h)
    para(tf, text, size, NAVY, bold=True, first=True)


def lead(slide, text, y=1.0, h=0.48, size=15):
    tf = box(slide, 0.48, y, 12.3, h)
    para(tf, text, size, INK, first=True, spacing=1.05)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rows(slide, items, y, h, x=0.48, w=12.36, label_w=3.15):
    for i, (left, right) in enumerate(items):
        fill = WHITE if i % 2 == 0 else PALE
        rect(slide, x, y, w, h, fill)
        tf = box(slide, x + 0.16, y, label_w, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, left, 13, NAVY, bold=True, first=True)
        tf = box(slide, x + 0.2 + label_w, y, w - label_w - 0.36, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, right, 13, INK, first=True)
        y += h


def cards(slide, items, y, cols, card_w, card_h, gap_x, gap_y, x0=0.48):
    for i, (num, head, body) in enumerate(items):
        col, row = i % cols, i // cols
        x = x0 + col * (card_w + gap_x)
        yy = y + row * (card_h + gap_y)
        rect(slide, x, yy, card_w, card_h, WHITE, LINE, rounded=True, radius=0.1)
        tf = box(slide, x + 0.16, yy + 0.1, 0.42, 0.32)
        para(tf, num, 14, COPPER, bold=True, first=True)
        tf = box(slide, x + 0.52, yy + 0.1, card_w - 0.7, 0.32)
        para(tf, head, 14, NAVY, bold=True, first=True)
        tf = box(slide, x + 0.16, yy + 0.46, card_w - 0.32, card_h - 0.56)
        para(tf, body, 12, INK, first=True, spacing=1.05)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "How Reddit tasks get published"
    prs.core_properties.subject = "Pitch for manager and CEO"
    slide_title(prs)
    slide_order(prs)
    slide_overview(prs)
    slide_what(prs)
    slide_problem(prs)
    slide_path(prs)
    slide_gates(prs)
    slide_join(prs)
    slide_rules(prs)
    slide_comment(prs)
    slide_post(prs)
    slide_dqn(prs)
    slide_rl_decides(prs)
    slide_rl_learns(prs)
    slide_random(prs)
    slide_benefits(prs)
    slide_results(prs)
    slide_ask(prs)
    slide_next_model(prs)
    slide_next_project(prs)
    prs.save(OUT)


def slide_title(prs):
    s = new_slide(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, 0.1, H, COPPER)
    tf = box(s, 0.62, 1.38, 11.5, 0.3)
    para(tf, "INTERNAL    ·    PITCH TO MANAGER AND CEO", 13, WARM, bold=True, first=True, spc=140)
    tf = box(s, 0.62, 1.88, 11.8, 1.7)
    para(tf, "How Reddit tasks", 40, WHITE, bold=True, first=True, spacing=0.95)
    para(tf, "get published", 40, WHITE, bold=True, spacing=0.95)
    tf = box(s, 0.62, 4.0, 11.2, 0.95)
    para(
        tf,
        "A general overview of the project first. Then the sitting. Then the model that keeps what Reddit left up.",
        18,
        SOFT,
        first=True,
        spacing=1.1,
    )
    labels = [("01", "Overview"), ("02", "Workflow"), ("03", "RL and DQN"), ("04", "Randomness")]
    x = 0.62
    for num, name in labels:
        rect(s, x, 5.45, 2.9, 0.58, "1B3A58", rounded=True, radius=0.14)
        tf = box(s, x + 0.14, 5.54, 2.62, 0.4, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = num + "  "
        style_run(r1, 13, COPPER, True)
        r2 = p.add_run()
        r2.text = name
        style_run(r2, 14, WHITE, True)
        x += 3.05
    tf = box(s, 0.62, 6.85, 4, 0.28)
    para(tf, "Klarivo", 14, WHITE, bold=True, first=True)
    notes(s, "Open with the general overview. Do not start with the model or the pattern.")


def slide_order(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "THE PITCH")
    title(s, "Four parts, in this order")
    lead(s, "General overview first, so the project is clear. Then the workflow. Then the model. Then the randomness.")
    items = [
        ("01", "Overview", "What the project is, what a sitting covers, and why posting by hand is costing us."),
        ("02", "Workflow", "How one sitting runs: account, rules, join, comment, and post."),
        ("03", "RL and DQN", "What the learning model decides, what it looks at, and how a result teaches it."),
        ("04", "Randomness", "How each sitting gets a new activity pattern, and why a used pattern is not repeated."),
    ]
    positions = [(0.48, 1.85), (6.7, 1.85), (0.48, 4.25), (6.7, 4.25)]
    for (num, head, body), (x, y) in zip(items, positions):
        rect(s, x, y, 6.1, 2.15, WHITE, LINE, rounded=True, radius=0.08)
        rect(s, x, y, 0.08, 2.15, COPPER)
        tf = box(s, x + 0.32, y + 0.22, 1.0, 0.4)
        para(tf, num, 18, COPPER, bold=True, first=True)
        tf = box(s, x + 1.3, y + 0.22, 4.5, 0.4)
        para(tf, head, 20, NAVY, bold=True, first=True)
        tf = box(s, x + 0.32, y + 0.85, 5.5, 1.0)
        para(tf, body, 15, INK, first=True, spacing=1.08)
    footer(s, 2)
    notes(s, "Stay in this order. The next slide is the general overview of the whole project.")


def slide_overview(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "01    ·    OVERVIEW")
    title(s, "General overview of the project")
    lead(
        s,
        "This is Klarivo’s Reddit publishing runner. Work comes from our sheets. Accounts that are already signed in do the activity. We record whether Reddit kept it.",
        y=0.98,
        h=0.52,
    )
    rect(s, 0.48, 1.58, 12.36, 0.78, NAVY, rounded=True, radius=0.08)
    tf = box(s, 0.7, 1.7, 11.95, 0.54, anchor=MSO_ANCHOR.MIDDLE)
    para(
        tf,
        "Sheets in  →  own browser  →  5–10 minutes on Reddit  →  comment or post  →  live, removed, or banned.",
        15,
        WHITE,
        bold=True,
        first=True,
    )
    pillars = [
        ("What it is", "A Python runner on AdsPower Chrome profiles. Each Reddit account stays in its own logged-in browser. The tool never types a password."),
        ("What it does", "Opens the account, browses Home, reads the community rules, joins, types a comment or a text post from the sheet, then checks if it stayed."),
        ("What it does not", "It does not create Reddit accounts, buy karma, run ads, or replace the sheets. People still prepare the comment and the post."),
        ("How it is run", "Up to 8 accounts at once. Two comments and one post per account every 48 hours. A banned or suspended account is stopped and named."),
    ]
    positions = [(0.48, 2.54), (6.7, 2.54), (0.48, 4.62), (6.7, 4.62)]
    for (head, body), (x, y) in zip(pillars, positions):
        rect(s, x, y, 6.1, 1.92, WHITE, LINE, rounded=True, radius=0.08)
        rect(s, x, y, 0.08, 1.92, COPPER)
        tf = box(s, x + 0.28, y + 0.16, 5.6, 0.38)
        para(tf, head, 16, NAVY, bold=True, first=True)
        tf = box(s, x + 0.28, y + 0.58, 5.6, 1.16)
        para(tf, body, 13, INK, first=True, spacing=1.08)
    footer(s, 3)
    notes(
        s,
        "This is the general overview. One sentence, the pipeline, then the four boxes. Do not walk the eight sitting steps here.",
    )


def slide_what(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "01    ·    OVERVIEW")
    title(s, "The six parts of the project")
    lead(s, "Same runner, six pieces. The sitting, the model, and the random pattern are the next sections.")
    items = [
        ("The task", "A comment link or a post row is prepared on the sheet. Unused sheet rows go first."),
        ("The account", "Each account stays in its own browser, already logged in. No password is typed."),
        ("The community", "It opens the community, reads the public rules, then joins. One new join per sitting."),
        ("The comment", "It types the comment on the post. It counts only if the comment is visible on the thread."),
        ("The post", "One text post when karma is 10 or more. The sheet title and body are not rewritten."),
        ("The model", "A learning model scores each choice from what Reddit left up and what it removed."),
    ]
    rows(s, items, 1.62, 0.84, label_w=2.35)
    footer(s, 4)
    notes(s, "Name the six. The workflow comes next. The model and the random pattern each get their own section after that.")


def slide_problem(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "01    ·    OVERVIEW")
    title(s, "What is going wrong today")
    lead(s, "A person takes the task and puts it on Reddit by hand. A lot of that work never counts.")
    cards_data = [
        ("Accounts get banned", "We lose the account and every task that was on it."),
        ("Comments get removed", "The reply never stays on the thread."),
        ("Posts get removed", "The post never counts for the client."),
        ("The task never counts", "We paid for the account and the work. The client gets nothing."),
    ]
    positions = [(0.48, 2.05), (6.7, 2.05), (0.48, 4.35), (6.7, 4.35)]
    for (head, body), (x, y) in zip(cards_data, positions):
        rect(s, x, y, 6.1, 2.05, WHITE, LINE, rounded=True, radius=0.08)
        rect(s, x, y, 0.08, 2.05, COPPER)
        tf = box(s, x + 0.32, y + 0.32, 5.5, 0.45)
        para(tf, head, 18, NAVY, bold=True, first=True)
        tf = box(s, x + 0.32, y + 0.9, 5.5, 0.8)
        para(tf, body, 15, INK, first=True)
    footer(s, 5)
    notes(s, "The cost is banned accounts, removed work, and a client who gets nothing from work we already paid for.")


def slide_path(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "How one sitting runs")
    lead(s, "One account. Already signed in. About 5 to 10 minutes. This is the order.")
    steps = [
        ("1", "Open", "Its own browser. Several accounts can run at once. The rest wait. Two accounts do not open the same community together."),
        ("2", "Read the account", "Username, karma, and age. A banned or suspended account stops here and is named."),
        ("3", "Browse Home", "Home first, then a community, then Home again. Most of the sitting stays on Home."),
        ("4", "Rules, then join", "Open the rules, read them, then join if it is not already in. One new join this sitting."),
        ("5", "Comment or read", "The model chooses, from those rules. The comment is typed on the post, or the thread is left."),
        ("6", "Post if allowed", "One text post when karma is 10 or more and the 48-hour slot is free."),
        ("7", "Check the result", "A comment counts only if it is visible. Older comments are scored again on the next run."),
        ("8", "Save", "Close the browser. Save what happened, so the next sitting starts from it."),
    ]
    cards(s, steps, 1.62, 2, 6.22, 1.22, 0.16, 0.12)
    footer(s, 6)
    notes(s, "Walk 1 to 8. The model is section 3. The fresh pattern on each sitting is section 4.")


def slide_gates(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "Warm-up, then the caps")
    lead(s, "A new account is not handed a client task. The caps stay on after warm-up.")
    items = [
        ("Under 7 days", "Browse only. No comment and no post."),
        ("Karma under 5", "Browse only. No comment and no post."),
        ("Karma 5 to 9", "Comments can run. Posts stay locked."),
        ("Karma 10", "One text post per 48 hours. A community can still ask for more."),
        ("48-hour budget", "2 comments and 1 post per account. Sheet work and general work share that budget."),
        ("Same thread once", "One comment per post URL, across every account. A used link is archived."),
        ("Already banned", "The run stops. The report names the account. It is not used."),
    ]
    rows(s, items, 1.58, 0.74, label_w=2.7)
    footer(s, 7)
    notes(s, "Karma under 5 is browse only. Posts wait for karma 10. Two comments and one post per 48 hours.")


def slide_join(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "Which community, and the join")
    lead(s, "It does not walk the same communities every sitting.")
    items = [
        ("From the sheet", "Enabled rows in the community sheet. Those are the communities for the work."),
        ("Not the last ones", "It prefers communities this account did not use last sitting."),
        ("Easier when new", "A low-karma account is pointed at easier communities first."),
        ("Never some", "A blocked list is never opened, even if a sheet names it."),
        ("Not together", "Two accounts in the same run do not open the same community at the same time."),
        ("One new join", "After warm-up, one Join click per sitting. If it is already in, it stays."),
        ("Join is scored", "A join that works is a plus for the model. A join that fails is a minus."),
    ]
    rows(s, items, 1.58, 0.74, label_w=2.55)
    footer(s, 8)
    notes(s, "One join per sitting, after the rules. The model is told whether the join worked. Extra random communities are in the randomness section.")


def slide_rules(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "It reads the rules before it speaks")
    lead(s, "Every community is opened on its rules page first. The comment is written from those rules.")
    items = [
        ("The rules page", "It opens the public rules for that community and reads them."),
        ("What it notes", "How many rules, how strict they are, and the flags. Jokes, age, karma."),
        ("Into the comment", "Those rules are part of the prompt that writes the comment."),
        ("Tone can be blocked", "Funny is dropped on a strict community. A gate on the account can also drop expert."),
        ("Clash is rewritten", "If the chosen tone fights the rules, the comment is rewritten plainer."),
        ("The model is marked", "A tone that fits is a small plus. A clash is a minus. Reading the rules is a plus."),
        ("Then it chooses", "Only after the rules: comment on a post, or only read."),
    ]
    rows(s, items, 1.58, 0.74, label_w=2.7)
    footer(s, 9)
    notes(s, "Rules are read, stored, put into the comment, and used to block a bad tone. Then the model picks comment or only read.")


def slide_comment(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "How a comment is added to a post")
    lead(s, "Sheet links go first. A comment is counted only when it is visible on the thread.")
    steps = [
        ("1", "Pick the post", "An unused link from the comment sheet, if there is one. Otherwise a new post in the community."),
        ("2", "Check the kind", "Only a question, help, a suggestion, or a review. Image, video, gallery, and GIF posts are skipped."),
        ("3", "Pick the tone", "The model chooses friendly, expert, funny, or neutral. Or it skips. Rules can remove a tone first."),
        ("4", "Write the words", "The exact text from the sheet, or one or two sentences about that post, using the rules."),
        ("5", "Read, then type", "It scrolls the thread, pauses, opens the comment box, and types with uneven keystrokes."),
        ("6", "Submit Comment", "It clicks Comment. It does not use the Post button. It stays on the thread afterwards."),
        ("7", "Count if visible", "If the text is on the thread, it counts. Blocked, filtered, or missing does not count."),
        ("8", "Record it", "A sheet link moves to the history sheet. The same URL is never commented twice. An optional edit is added 3 to 4 minutes later."),
    ]
    cards(s, steps, 1.58, 2, 6.22, 1.24, 0.16, 0.1)
    footer(s, 10)
    notes(s, "Sheet first, right kind of post, tone from the model, typed into the comment box, counted only if visible.")


def slide_post(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "02    ·    WORKFLOW")
    title(s, "How a post is added")
    lead(s, "Posts are the stricter action. Karma 10, one per 48 hours, and the community has to fit a text post.")
    items = [
        ("Sheet first", "If a post row is waiting, that title and body are used exactly. They are not rewritten."),
        ("Karma 10", "Below 10, the account can comment and cannot post."),
        ("One per 48 hours", "One live post. If the slot was already used, it waits."),
        ("If the sheet is empty", "It can write one post that fits the community, after it has read recent posts there."),
        ("Text posts only", "It checks that text posts are normal there. An image-only community is skipped."),
        ("Flair, then Post", "It picks the community tag from the post, then clicks Post."),
        ("If it fails", "The model ranks the next community. A removed post is a strong minus."),
    ]
    rows(s, items, 1.58, 0.74, label_w=2.7)
    footer(s, 11)
    notes(s, "Posts wait for karma 10. That is why most sittings still skip the post.")


def slide_dqn(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "03    ·    RL AND DQN")
    title(s, "What RL and the DQN are")
    lead(s, "Reinforcement learning means the result of an action changes the next choice. The DQN is the small network that scores those choices.")
    items = [
        ("Reinforcement learning", "It tries an action. Reddit answers. That answer is a reward, and the next choice uses it."),
        ("The DQN", "Deep Q-Network. A small neural net. It gives every allowed action a score. A higher score means a better expected result."),
        ("How a choice is made", "It compares the scores. A strong score is chosen more often. A weak score is rare."),
        ("What it remembers", "Past outcomes are stored and trained on again, so a later sitting starts from what already happened."),
        ("A thin random slice", "A small share of choices stays random, so a tone it has barely tried can still be picked. That share shrinks as results come in."),
        ("What it does not do", "It does not pick the browser, the scroll, or the sitting pattern. That pattern is the next section."),
    ]
    rows(s, items, 1.62, 0.84, label_w=3.15)
    footer(s, 12)
    notes(s, "RL is learning from the result. DQN is the network that scores the actions. Say clearly that the browse pattern is not the model.")


def slide_rl_decides(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "03    ·    RL AND DQN")
    title(s, "What the model decides")
    lead(s, "Four decisions. The score for each one depends on the account, the post, the rules, and the moment.")
    left = [
        ("Comment or read", "After the rules are read. Comment, or only lurk."),
        ("The tone", "Friendly, expert, funny, or neutral."),
        ("Skip", "Leave the thread when this kind of post has been getting removals."),
        ("Next community", "If the first community refuses the post, which one to try."),
    ]
    right = [
        ("Karma and age", "A new account is not treated like an old one."),
        ("The post", "Mood, length, and whether it is a question, help, a suggestion, or a review."),
        ("The rules", "Count, strictness, flags, and whether they were read."),
        ("The moment", "Hour, weekday, comments already left, and recent success."),
    ]
    tf = box(s, 0.48, 1.58, 6, 0.28)
    para(tf, "It chooses", 14, COPPER, bold=True, first=True)
    tf = box(s, 6.9, 1.58, 6, 0.28)
    para(tf, "It looks at", 14, COPPER, bold=True, first=True)
    y = 1.96
    for (lh, lb), (rh, rb) in zip(left, right):
        rect(s, 0.48, y, 6.15, 1.12, WHITE, LINE, rounded=True, radius=0.1)
        rect(s, 6.82, y, 6.05, 1.12, WHITE, LINE, rounded=True, radius=0.1)
        tf = box(s, 0.66, y + 0.12, 5.8, 0.32)
        para(tf, lh, 15, NAVY, bold=True, first=True)
        tf = box(s, 0.66, y + 0.48, 5.8, 0.5)
        para(tf, lb, 13, INK, first=True)
        tf = box(s, 7.0, y + 0.12, 5.7, 0.32)
        para(tf, rh, 15, NAVY, bold=True, first=True)
        tf = box(s, 7.0, y + 0.48, 5.7, 0.5)
        para(tf, rb, 13, INK, first=True)
        y += 1.22
    footer(s, 13)
    notes(s, "Left is the decision. Right is the state the network scores against.")


def slide_rl_learns(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "03    ·    RL AND DQN")
    title(s, "How it learns from the activity")
    lead(s, "Every outcome is stored. The network is trained on those outcomes, and the next run starts from them.")
    steps = [
        ("1", "Right away", "Rules read. Join worked or failed. Comment visible or not. Tone fit the rules or clashed."),
        ("2", "Later", "On the next run, a comment older than 20 minutes is checked again. Still live, the score, filtered, or removed."),
        ("3", "The score trains it", "Upvotes raise that choice. Filtered marks it to avoid. Removed is a strong minus, and that choice is dropped."),
        ("4", "It keeps the history", "Unknown after 3 days is closed. The past stays, so learning is not restarted each run."),
    ]
    y = 1.58
    for num, head, body in steps:
        rect(s, 0.48, y, 12.36, 0.92, WHITE, LINE, rounded=True, radius=0.1)
        tf = box(s, 0.66, y + 0.22, 0.4, 0.45, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, num, 16, COPPER, bold=True, first=True)
        tf = box(s, 1.15, y + 0.1, 2.4, 0.7, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, head, 14, NAVY, bold=True, first=True)
        tf = box(s, 3.6, y + 0.1, 9.0, 0.72, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, body, 13, INK, first=True)
        y += 1.0
    scores = [("Visible", "Good"), ("Upvotes", "Better"), ("Filtered", "Avoid"), ("Removed", "Stop it")]
    x = 0.48
    for label, meaning in scores:
        rect(s, x, 5.68, 3.02, 1.15, NAVY, rounded=True, radius=0.1)
        tf = box(s, x + 0.16, 5.8, 2.7, 0.36)
        para(tf, label, 13, WARM, bold=True, first=True)
        tf = box(s, x + 0.16, 6.16, 2.7, 0.42)
        para(tf, meaning, 16, WHITE, bold=True, first=True)
        x += 3.18
    footer(s, 14)
    notes(s, "Immediate reward, then a delayed check after 20 minutes. Removed choices get dropped. The history is kept between runs.")


def slide_random(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "04    ·    RANDOMNESS")
    title(s, "How randomness is added")
    lead(s, "This is the browse pattern, not the model. Each account rolls a fresh pattern, and a used pattern is not rolled again.")
    items = [
        ("Rolled at the start", "Before the browsing, the sitting rolls a new pattern for that account."),
        ("Six personas", "Slow lurker, curious, restless, steady, skimmer, or methodical. Each one changes the pace."),
        ("Timings move", "Sitting length, Home time, pauses, scroll, upvote chance, and typing speed are all rolled again."),
        ("Home first", "About half to three quarters of the sitting stays on Home. Communities are short visits."),
        ("Home between", "It usually goes back to Home between communities, instead of hopping community to community."),
        ("Extra communities", "1 to 4 extra communities are explored, on top of the sheet. About 30% of the list is exploration."),
        ("One search", "At most one short search. A search sentence that account already used is not typed again."),
        ("Not repeated", "The fingerprint is saved. The next sitting for that account has to be a different pattern."),
    ]
    rows(s, items, 1.55, 0.66, label_w=2.7)
    footer(s, 15)
    notes(s, "Randomness is the pattern: persona, timings, Home, extra communities, one search. The fingerprint blocks a repeat.")


def slide_benefits(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "FOR THE COMPANY")
    title(s, "What the company gets")
    items = [
        ("01", "People stop posting by hand", "They prepare the comment and the post. The tool does the Reddit activity."),
        ("02", "The same mistakes stop repeating", "The model drops a tone or a thread Reddit has already taken down."),
        ("03", "Accounts last longer", "A week of reading, karma gates, one join, and a 48-hour cap. A banned account is stopped."),
        ("04", "The work does not look like one script", "Each sitting is a new pattern. A used pattern is not repeated."),
        ("05", "A number next to the client", "Assigned, published, live, removed, banned. After every run."),
    ]
    y = 1.22
    for num, head, body in items:
        rect(s, 0.48, y, 12.36, 1.05, WHITE, LINE, rounded=True, radius=0.1)
        rect(s, 0.48, y, 0.08, 1.05, COPPER)
        tf = box(s, 0.74, y + 0.14, 0.5, 0.75, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, num, 14, COPPER, bold=True, first=True)
        tf = box(s, 1.4, y + 0.12, 11.1, 0.38)
        para(tf, head, 16, NAVY, bold=True, first=True)
        tf = box(s, 1.4, y + 0.52, 11.1, 0.4)
        para(tf, body, 14, INK, first=True)
        y += 1.14
    footer(s, 16)
    notes(s, "Tie each benefit to a section: hand posting, the model, warm-up, the random pattern, and the report.")


def slide_results(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "FOR THE COMPANY")
    title(s, "What we measure after a run")
    lead(s, "The report is the number next to the client. A comment counts only if it is still on the thread.")
    items = [
        ("Assigned", "Tasks taken from the sheets for this run."),
        ("Published", "Comments and posts the tool submitted."),
        ("Live", "Still visible on the thread when we check."),
        ("Filtered", "Submitted, then held by Reddit."),
        ("Removed", "Taken down. That choice is trained against."),
        ("Banned", "The account is stopped and named. It is not used again."),
    ]
    rows(s, items, 1.62, 0.82, label_w=2.35)
    footer(s, 17)
    notes(s, "Walk assigned through banned. Live is the number that matters for the client.")


def slide_ask(prs):
    s = new_slide(prs)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, 0.1, H, COPPER)
    tf = box(s, 0.58, 0.38, 10, 0.28)
    para(tf, "THE ASK", 12, WARM, bold=True, first=True, spc=140)
    tf = box(s, 0.58, 0.78, 12.1, 1.15)
    para(tf, "Run it on real accounts for 30 days.", 32, WHITE, bold=True, first=True)
    para(tf, "Same rules. The model keeps learning. The pattern keeps changing.", 18, SOFT, before=8)
    blocks = [
        ("30 days", "One set of accounts."),
        ("Sheets in", "Comments, posts, and communities come from the sheets."),
        ("Rules on", "7-day warm-up, karma 5 to comment, karma 10 to post, 48-hour caps."),
        ("Day 30", "Ban rate, removal rate, and tasks that stayed live, against today."),
    ]
    x = 0.58
    for head, body in blocks:
        rect(s, x, 3.2, 3.0, 1.9, "1B3A58", rounded=True, radius=0.08)
        tf = box(s, x + 0.16, 3.38, 2.68, 0.4)
        para(tf, head, 16, WARM, bold=True, first=True)
        tf = box(s, x + 0.16, 3.88, 2.68, 1.0)
        para(tf, body, 14, WHITE, first=True, spacing=1.05)
        x += 3.15
    tf = box(s, 0.58, 5.4, 12.1, 1.15)
    para(
        tf,
        "If those three move the right way, this becomes how these Reddit tasks get published.",
        16,
        SOFT,
        first=True,
    )
    para(
        tf,
        "After every run: live, filtered, removed, or the account banned.",
        16,
        SOFT,
        before=8,
    )
    tf = box(s, 0.58, 6.85, 4, 0.28)
    para(tf, "Klarivo", 14, WHITE, bold=True, first=True)
    tf = box(s, 9.3, 6.85, 3.4, 0.28)
    para(tf, f"18   /   {TOTAL}", 12, "8FA0B0", align=PP_ALIGN.RIGHT, first=True)
    notes(s, "Ask for 30 days on one set of accounts. The last two slides are the future plan for the model, then for the project.")


def slide_next_model(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "NEXT    ·    THE MODEL")
    title(s, "Future plan for the model")
    lead(s, "The DQN already scores comment vs lurk and the tone. Next is to make those scores match what Reddit actually keeps.")
    items = [
        ("Kept-work target", "Raise the share of comments and posts that stay live toward 80%. Drop a tone or thread Reddit already took down."),
        ("Train on the result", "Visible and upvotes raise that choice. Filtered means avoid. Removed is a strong minus, and that choice is dropped."),
        ("Lurk when weak", "If the score says the comment will not stay, the account only reads. A comment is not forced."),
        ("Community memory", "If a community restricts the post or the join fails, it is not chosen again for that kind of work."),
        ("Same four decisions", "Comment or lurk, tone, skip, and next community keep sharing one set of scores."),
        ("History stays on", "The next run starts from past outcomes. A thin random slice shrinks as real results come in."),
    ]
    rows(s, items, 1.58, 0.84, label_w=3.15)
    footer(s, 19)
    notes(s, "This is the model plan only. Target 80% kept work, train on live vs removed, lurk when weak, remember failed communities, keep history.")


def slide_next_project(prs):
    s = new_slide(prs)
    content_bg(s)
    eyebrow(s, "NEXT    ·    THE PROJECT")
    title(s, "Future plan for the project")
    lead(s, "The runner is built. Next is a live 30-day run, then scale only what the numbers support.")
    items = [
        ("30-day live run", "One set of real accounts. Same sheets, same caps, 8 at a time, 5 to 10 minutes. The model keeps learning."),
        ("The report", "After every run: assigned, published, live, filtered, removed, banned. Live is the number that counts."),
        ("Skip what will not count", "Do not post into a community that refuses. Do not spend the sitting on a long search or a join that will not count."),
        ("Pattern stays fresh", "Each sitting rolls a new browse pattern. A used pattern is not repeated for that account."),
        ("Scale only what worked", "Add accounts only if the ban rate stays down. Expand communities only where work stays live. Tune the 48-hour caps from the 30-day numbers."),
        ("Hand posting stops", "People prepare the comment and the post on the sheets. The tool publishes. This becomes how the work is done."),
    ]
    rows(s, items, 1.58, 0.84, label_w=3.15)
    footer(s, 20)
    notes(s, "Last slide. Project plan: 30-day run, the report, skip wasted work, keep the pattern fresh, scale only if bans and live work move the right way.")


if __name__ == "__main__":
    build()
