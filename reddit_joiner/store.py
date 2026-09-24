#!/usr/bin/env python3
"""SQLite activity log and live_posts.xlsx helpers."""

from __future__ import annotations

import hashlib
import os
import sqlite3
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

from reddit_joiner.paths import DB_PATH as _DB_PATH, LIVE_POSTS_XLSX as _LIVE_POSTS_XLSX

DB_PATH = str(_DB_PATH)
LIVE_POSTS_XLSX = str(_LIVE_POSTS_XLSX)


def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=60)
    conn.row_factory = sqlite3.Row
    return conn


def _ensure_column(conn: sqlite3.Connection, table: str, column: str, ddl: str) -> None:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    names = {str(row[1]) for row in rows}
    if column not in names:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {ddl}")


def init_db() -> None:
    with _connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS ai_comments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                account TEXT,
                user_id TEXT,
                post_url TEXT,
                post_title TEXT,
                sentiment REAL,
                comment TEXT,
                method TEXT
            );
            CREATE TABLE IF NOT EXISTS posts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                account TEXT,
                user_id TEXT,
                subreddit TEXT,
                title TEXT,
                url TEXT,
                status TEXT,
                reason TEXT,
                score INTEGER
            );
            CREATE TABLE IF NOT EXISTS post_attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fingerprint TEXT,
                post_title TEXT,
                post_body TEXT,
                original_subreddit TEXT,
                attempted_subreddit TEXT,
                account_id TEXT,
                attempt_number INTEGER,
                status TEXT,
                post_url TEXT,
                removed_reason TEXT,
                timestamp TEXT
            );
            CREATE TABLE IF NOT EXISTS account_comment_stats (
                account_id TEXT PRIMARY KEY,
                total_comments INTEGER DEFAULT 0,
                total_upvotes INTEGER DEFAULT 0,
                last_comment TEXT
            );
            """
        )
        _ensure_column(conn, "ai_comments", "subreddit", "TEXT")
        _ensure_column(conn, "ai_comments", "post_body", "TEXT")
        _ensure_column(conn, "ai_comments", "comment_id", "TEXT")
        _ensure_column(conn, "ai_comments", "upvotes", "INTEGER DEFAULT 0")
        _ensure_column(conn, "ai_comments", "tone", "TEXT")
        _ensure_column(conn, "ai_comments", "status", "TEXT")
        _ensure_column(conn, "posts", "fingerprint", "TEXT")
        _ensure_column(conn, "posts", "attempt_number", "INTEGER")


def log_ai_comment(
    *,
    account: str,
    user_id: str,
    post_url: str,
    post_title: str,
    sentiment: float,
    comment: str,
    method: str,
    subreddit: str = "",
    post_body: str = "",
    comment_id: str = "",
    tone: str = "",
) -> None:
    init_db()
    stamp = datetime.now().isoformat(timespec="seconds")
    with _connect() as conn:
        conn.execute(
            """
            INSERT INTO ai_comments
            (created_at, account, user_id, post_url, post_title, sentiment, comment, method,
             subreddit, post_body, comment_id, upvotes, tone)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?)
            """,
            (
                stamp,
                account,
                user_id,
                post_url,
                post_title,
                sentiment,
                comment,
                method,
                subreddit,
                post_body,
                comment_id,
                tone,
            ),
        )
        conn.execute(
            """
            INSERT INTO account_comment_stats (account_id, total_comments, total_upvotes, last_comment)
            VALUES (?, 1, 0, ?)
            ON CONFLICT(account_id) DO UPDATE SET
                total_comments = total_comments + 1,
                last_comment = excluded.last_comment
            """,
            (user_id or account, stamp),
        )


def log_post_attempt(
    *,
    account: str,
    user_id: str,
    subreddit: str,
    title: str,
    url: str,
    status: str,
    reason: str = "",
    score: int = 0,
    body: str = "",
) -> None:
    init_db()
    # The fingerprint has to be stored here, not just on post_attempts.
    # PostRetryManager.is_live() looks for a LIVE row by fingerprint, so leaving
    # it NULL made every successful post invisible to the duplicate check and
    # let the same title and body go out again from another account.
    fingerprint = post_fingerprint(title, body)
    with _connect() as conn:
        conn.execute(
            """
            INSERT INTO posts
            (created_at, account, user_id, subreddit, title, url, status, reason, score,
             fingerprint)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now().isoformat(timespec="seconds"),
                account,
                user_id,
                subreddit,
                title,
                url,
                status,
                reason,
                score,
                fingerprint,
            ),
        )


def append_live_post(
    *,
    account: str,
    subreddit: str,
    title: str,
    url: str,
    score: int = 0,
    status: str = "LIVE",
) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return
    headers = ["posted_at", "account", "subreddit", "title", "url", "score", "status"]
    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        account,
        subreddit,
        title,
        url,
        score,
        status,
    ]
    if os.path.isfile(LIVE_POSTS_XLSX):
        book = load_workbook(LIVE_POSTS_XLSX)
        sheet = book.active
        if sheet.max_row == 1 and not any(
            (sheet.cell(1, col).value or "") for col in range(1, len(headers) + 1)
        ):
            sheet.append(headers)
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "live_posts"
        sheet.append(headers)
    sheet.append(row)
    book.save(LIVE_POSTS_XLSX)


def recent_live_posts(limit: int = 20) -> List[Dict[str, Any]]:
    init_db()
    with _connect() as conn:
        rows = conn.execute(
            "SELECT * FROM posts WHERE status = 'LIVE' ORDER BY id DESC LIMIT ?",
            (limit,),
        ).fetchall()
    return [dict(row) for row in rows]


def _link_key(url: str) -> str:
    return (url or "").split("?")[0].split("#")[0].rstrip("/").lower()


def post_fingerprint(title: str, body: str) -> str:
    blob = f"{(title or '').strip()}\n{(body or '').strip()}".encode("utf-8")
    return hashlib.sha256(blob).hexdigest()[:20]


def comments_today_for_account(user_id: str) -> int:
    init_db()
    start = datetime.now().strftime("%Y-%m-%d")
    with _connect() as conn:
        row = conn.execute(
            """
            SELECT COUNT(*) AS n FROM ai_comments
            WHERE (user_id = ? OR account = ?) AND created_at >= ?
            """,
            (user_id, user_id, start),
        ).fetchone()
    return int(row["n"] if row else 0)


def seconds_since_last_ai_comment(user_id: str) -> Optional[float]:
    init_db()
    with _connect() as conn:
        row = conn.execute(
            """
            SELECT created_at FROM ai_comments
            WHERE user_id = ? OR account = ?
            ORDER BY id DESC LIMIT 1
            """,
            (user_id, user_id),
        ).fetchone()
    if not row or not row["created_at"]:
        return None
    try:
        stamp = datetime.fromisoformat(str(row["created_at"]))
    except ValueError:
        return None
    return (datetime.now() - stamp).total_seconds()


def attach_ai_comment_permalink(post_url: str, comment_url: str, comment_id: str = "") -> None:
    """Remember the comment permalink / id on the latest matching ai_comments row."""
    key = _link_key(post_url)
    cid = (comment_id or "").strip()
    if not cid and comment_url:
        parts = [part for part in (comment_url or "").split("/") if part]
        if parts:
            last = parts[-1]
            if 5 <= len(last) <= 12 and last.replace("_", "").isalnum():
                cid = last.replace("t1_", "")
    if not key and not cid:
        return
    init_db()
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id, post_url, comment_id FROM ai_comments ORDER BY id DESC LIMIT 80"
        ).fetchall()
        target = None
        for row in rows:
            if key and _link_key(str(row["post_url"] or "")) == key:
                target = row["id"]
                break
            existing = str(row["comment_id"] or "").strip()
            if cid and existing == cid:
                target = row["id"]
                break
        if target is None:
            return
        conn.execute(
            "UPDATE ai_comments SET comment_id = COALESCE(NULLIF(comment_id, ''), ?) WHERE id = ?",
            (cid, target),
        )


def unscored_ai_comments(limit: int = 40) -> List[Dict[str, Any]]:
    """Recent comments that have not been scored live/removed yet."""
    init_db()
    with _connect() as conn:
        rows = conn.execute(
            """
            SELECT created_at, account, user_id, post_url, post_title, sentiment,
                   comment, method, subreddit, post_body, comment_id, tone, status
            FROM ai_comments
            WHERE IFNULL(status, '') = ''
            ORDER BY id DESC
            LIMIT ?
            """,
            (max(1, int(limit)),),
        ).fetchall()
    return [dict(row) for row in rows]


def update_ai_comment_outcome(url: str, *, score: int = 0, status: str = "") -> None:
    """Store the latest live/removed/score check for a posted comment."""
    key = _link_key(url)
    if not key:
        return
    init_db()
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id, post_url, comment_id, upvotes, status FROM ai_comments "
            "ORDER BY id DESC LIMIT 400"
        ).fetchall()
        # Match on the comment id first and only by whole path segment. The old
        # order checked post_url first, so an outcome for one account's comment
        # could be stamped onto whichever row most recently touched that post,
        # and a substring test let a short id match inside an unrelated path.
        target = None
        previous_score = 0
        segments = [part for part in key.split("/") if part]
        for row in rows:
            comment_id = str(row["comment_id"] or "").strip().lower()
            if comment_id and comment_id in segments:
                target = row["id"]
                previous_score = int(row["upvotes"] or 0)
                break
        if target is None:
            # Fall back to the post permalink, but only when exactly one row
            # references that post. If several accounts commented on it, the
            # bare post URL cannot say which one is being graded, and guessing
            # is what wrote outcomes onto the wrong rows.
            matches = [
                row for row in rows if _link_key(str(row["post_url"] or "")) == key
            ]
            if len(matches) != 1:
                return
            target = matches[0]["id"]
            previous_score = int(matches[0]["upvotes"] or 0)
        conn.execute(
            "UPDATE ai_comments SET upvotes = ?, status = ? WHERE id = ?",
            (int(score or 0), str(status or ""), target),
        )
        # Apply only the change since the last check. Re-checking a live comment
        # used to add its full score again on every pass.
        delta = int(score or 0) - int(previous_score or 0)
        if status == "live" and delta:
            conn.execute(
                """
                UPDATE account_comment_stats
                SET total_upvotes = MAX(0, COALESCE(total_upvotes, 0) + ?)
                WHERE account_id IN (
                    SELECT COALESCE(user_id, account) FROM ai_comments WHERE id = ?
                )
                """,
                (delta, target),
            )


def already_commented_url(user_id: str, url: str) -> bool:
    key = _link_key(url)
    if not key:
        return False
    init_db()
    with _connect() as conn:
        rows = conn.execute(
            "SELECT post_url FROM ai_comments WHERE user_id = ? OR account = ?",
            (user_id, user_id),
        ).fetchall()
    return any(_link_key(str(row["post_url"] or "")) == key for row in rows)


def already_commented_url_any(url: str) -> bool:
    """True if any AdsPower account already left a comment on this post."""
    key = _link_key(url)
    if not key:
        return False
    init_db()
    with _connect() as conn:
        rows = conn.execute("SELECT post_url FROM ai_comments").fetchall()
    return any(_link_key(str(row["post_url"] or "")) == key for row in rows)


def comments_in_last_days(user_id: str, days: float) -> List[Dict[str, Any]]:
    if not user_id:
        return []
    init_db()
    cutoff = (datetime.now() - timedelta(days=float(days))).isoformat(timespec="seconds")
    with _connect() as conn:
        rows = conn.execute(
            """
            SELECT id, created_at, post_url, comment
            FROM ai_comments
            WHERE (user_id = ? OR account = ?) AND created_at >= ?
            ORDER BY id DESC
            """,
            (user_id, user_id, cutoff),
        ).fetchall()
    return [dict(row) for row in rows]


def can_ai_comment(
    user_id: str,
    *,
    max_per_day: int,
    min_delay_seconds: float,
) -> Tuple[bool, str]:
    if not user_id:
        return False, "no account"
    today = comments_today_for_account(user_id)
    if today >= max_per_day:
        return False, f"daily cap {today}/{max_per_day}"
    elapsed = seconds_since_last_ai_comment(user_id)
    if elapsed is not None and elapsed < min_delay_seconds:
        wait = min_delay_seconds - elapsed
        return False, f"wait {wait:.0f}s before next AI comment"
    return True, "ok"


class PostRetryManager:
    """Track spreadsheet post attempts and LIVE rows across accounts."""

    def __init__(self) -> None:
        init_db()

    def log_attempt(
        self,
        *,
        post_title: str,
        post_body: str,
        subreddit: str,
        account_id: str,
        attempt_num: int,
        status: str,
        url: str = "",
        reason: str = "",
        original_subreddit: str = "",
    ) -> None:
        init_db()
        fingerprint = post_fingerprint(post_title, post_body)
        with _connect() as conn:
            conn.execute(
                """
                INSERT INTO post_attempts
                (fingerprint, post_title, post_body, original_subreddit, attempted_subreddit,
                 account_id, attempt_number, status, post_url, removed_reason, timestamp)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    fingerprint,
                    post_title,
                    post_body,
                    original_subreddit or subreddit,
                    subreddit,
                    account_id,
                    attempt_num,
                    status,
                    url,
                    reason,
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )

    def get_attempts_for_post(self, post_title: str, post_body: str) -> List[Dict[str, Any]]:
        init_db()
        fingerprint = post_fingerprint(post_title, post_body)
        with _connect() as conn:
            rows = conn.execute(
                "SELECT * FROM post_attempts WHERE fingerprint = ? ORDER BY id",
                (fingerprint,),
            ).fetchall()
        return [dict(row) for row in rows]

    def failed_subreddits(self, post_title: str, post_body: str, account_id: str = "") -> List[str]:
        failed = []
        seen = set()
        for row in self.get_attempts_for_post(post_title, post_body):
            if account_id and str(row.get("account_id") or "") != account_id:
                continue
            status = str(row.get("status") or "").upper()
            if status in {"LIVE", "APPROVED"}:
                continue
            name = str(row.get("attempted_subreddit") or "").lower()
            if name and name not in seen:
                seen.add(name)
                failed.append(name)
        return failed

    def is_live(self, post_title: str, post_body: str) -> bool:
        for row in self.get_attempts_for_post(post_title, post_body):
            if str(row.get("status") or "").upper() in {"LIVE", "APPROVED"}:
                return True
        init_db()
        fingerprint = post_fingerprint(post_title, post_body)
        with _connect() as conn:
            row = conn.execute(
                "SELECT id FROM posts WHERE fingerprint = ? AND status = 'LIVE' LIMIT 1",
                (fingerprint,),
            ).fetchone()
        return bool(row)

    def accounts_tried(self, post_title: str, post_body: str) -> List[str]:
        seen = []
        keys = set()
        for row in self.get_attempts_for_post(post_title, post_body):
            account = str(row.get("account_id") or "")
            if account and account not in keys:
                keys.add(account)
                seen.append(account)
        return seen

    def should_retry_on_another_account(
        self,
        post_title: str,
        post_body: str,
        max_accounts: int = 2,
    ) -> bool:
        if self.is_live(post_title, post_body):
            return False
        return len(self.accounts_tried(post_title, post_body)) < max_accounts

    def save_live_post(
        self,
        *,
        post_title: str,
        post_body: str,
        subreddit: str,
        account_id: str,
        url: str,
        score: int = 0,
    ) -> None:
        append_live_post(
            account=account_id,
            subreddit=subreddit,
            title=post_title,
            url=url,
            score=score,
            status="LIVE",
        )
